"""Puts the UI figures exported from docs/wireframes into the FDS (Report_3.2_FDS_IELTSPath_v1.docx).

Usage:  python update_fds_figures.py [input.docx] [output.docx]   (defaults: the FDS next to this script, in place)

The draw.io files live in docs/wireframes/screens (one file per P-ID, shared with the Screen Design Spec). Build and
export them first with the Screen Design Spec scripts:
    ../rp_3.2_screen_design_spec/scripts/build_diagrams.py  then  export_diagrams.py
Safe to run again: in every screen of SCREENS the figures between the General Information table and
"Navigation Context" are replaced, the captions renumbered and the Mockup / Wireframe row rewritten. Image parts
no longer used are dropped from the package.
"""
import copy
import io
import os
import sys
import xml.etree.ElementTree as ET

import docx
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(HERE)
WIREFRAMES = os.path.join(DOCS, "wireframes")
sys.path.insert(0, os.path.join(DOCS, "rp_3.2_screen_design_spec", "scripts"))
from export_diagrams import export_name  # noqa: E402
from sds_screens import by_id  # noqa: E402

FDS = os.path.join(HERE, "Report_3.2_FDS_IELTSPath_v1.docx")
W_P, W_TBL = qn("w:p"), qn("w:tbl")
WIDTH_CM, MAX_HEIGHT_CM = 15.5, 21.0
EMBED_MAX_PX = 1600  # embedded copy only; wireframes/exports keeps the full-size PNG

# FDS section, screen name, sources: draw.io file stem in wireframes/screens, optionally "stem:Frame|Frame"
SCREENS = [
    ("2.1", "Home", ["P-10_Home", "P-10c_KeyActivation"]),
    ("2.2", "Login", ["P-01_SignIn"]),
    ("2.3", "Register", ["P-02_SignUp"]),
    ("2.4", "Forgot Password", ["P-03_ForgotPassword"]),
    ("2.5", "Reset Password", ["P-04_ResetPassword"]),
    ("2.7", "Overview", ["P-11_OverviewDashboard"]),
    ("2.8", "Course List", ["P-30_CourseList"]),
    ("2.9", "Placement Test",
     ["P-30a_Placement:Survey|Survey Summary|Test Hub|Exam — Reading / Listening|Exam — Writing|Exam — Speaking|Grading"]),
    ("2.10", "Placement Report", ["P-30a_Placement:Result Report"]),
    ("2.11", "Course Path", ["P-30b_TopicList"]),
    ("2.12", "Topic Detail", ["P-31_TopicDetail"]),
    ("2.13", "Lesson Player", ["P-32_LessonPlayer:Populated|Completed|Review Required"]),
    ("2.14", "Writing Feedback", ["P-32_LessonPlayer:Writing Feedback — Chưa đạt|Writing Feedback — Đã đạt"]),
    ("2.15", "Practice Set", ["P-33_LessonPractice"]),
    ("2.17", "Review Session", ["P-34_ReviewSession"]),
    ("2.19", "Test Attempt", ["P-35_TopicTest"]),
    ("2.20", "Test Result", ["P-36_TestResult"]),
    ("2.21", "Mock Tests", ["P-46_MockTestCatalog", "P-47_MockTestRunner"]),
    ("2.22", "Practice Test Catalog", ["P-41_PracticeCatalog", "P-41a_PracticeModeSelect"]),
    ("2.23", "Practice Workspaces",
     ["P-42_ReadingPracticeTest", "P-42a_CreateFlashcard", "P-42b_DictionaryLookup", "P-42c_SavedFlashcards",
      "P-42d_FloatingNotes", "P-43_ListeningPractice", "P-44_WritingPractice", "P-44a_WritingChartZoom",
      "P-45_SpeakingPractice"]),
    ("2.24", "Classroom", ["P-20_MyClassroom", "P-20a_LessonDetail"]),
    ("2.25", "Classroom Lesson Workspace", ["P-21_LessonWorkspace"]),
    ("2.35", "Flashcards & Decks", ["P-51_MyFlashcards", "P-51a_CreateSet", "P-51b_SetDetail"]),
    ("2.26", "Profile & Account", ["P-80_Account:Hồ sơ của bạn|Lịch sử đăng nhập", "SHELL_SharedLayout:User Menu"]),
    ("2.28", "Subscription & Points", ["P-80_Account:Lịch sử kích hoạt gói|Lịch sử giao dịch điểm"]),
    ("2.31", "Vocabulary", ["P-50_Vocabulary", "P-56_DictionaryHome", "P-57_WordDetail", "P-58_VocabularyMap",
                            "P-59_VocabularySet", "P-59a_SetStudy", "P-59b_TopicLessonPath", "P-59c_VocabularyLesson",
                            "P-59d_Flashcard"]),
    ("2.54", "Submission History", ["P-54_SubmissionHistory"]),
]
V07_DESCRIPTION = (
    "Aligned with the frontend at 6525a6d (main) and the backend at eedc6cd (no code change since bf7db56; the "
    "capstone documents moved to the docs repository): Overview follows the refreshed dashboard — Dashboard hover menu, hero "
    "with the current course and IELTS Journey, five study stats, the activity calendar with day and week selection, "
    "My Learning Orbit, Study Plan, SPACE Autonomy, AI suggestions and the Dữ liệu chi tiết tab; Placement Test waits "
    "for GET placement-result instead of the course list, so FG-17 is resolved; every Part 2 UI figure is now exported "
    "from the draw.io wireframes in docs/wireframes/screens (one file per P-ID, shared with the Screen Design Spec; "
    "clickable prototype docs/wireframes/IELTSPath_Prototype.drawio) in place of the application screenshots, with "
    "new drawings for Overview (P-11), Writing Feedback (P-32), Classroom (P-20, P-20a), Classroom Lesson Workspace "
    "(P-21) and Vocabulary (P-50) and figures added to every built screen; design-target wireframes added for Mock "
    "Tests (P-46 catalog, P-47 computer-based runner), Speaking Practice (P-45), the 3-column Writing Practice (P-44), "
    "Profile & Account and Subscription & Points (P-80); Submission History (P-54) is specified in Section 2.54. "
    "The Dashboard menu replaces its left sidebar; personal information is reached through Thông tin cá nhân in "
    "the user menu. The account tab is named Lịch sử kích hoạt gói and the dictionary remains Tra Anh-Việt.")


def frames(source):
    """(tab name, PNG path) of the tabs a source names, in file order."""
    stem, _, wanted = source.partition(":")
    path = os.path.join(WIREFRAMES, "screens", stem + ".drawio")
    tabs = [d.get("name") for d in ET.parse(path).getroot().findall("diagram")]
    keep = wanted.split("|") if wanted else None
    return [(t, os.path.join(WIREFRAMES, "exports", export_name(t) + ".png")) for t in tabs
            if keep is None or t.split(" / ", 1)[1] in keep]


def ptext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def set_p(p_el, text):
    """Keeps the first run's formatting, drops the other runs."""
    runs = p_el.findall(qn("w:r"))
    for r in runs[1:]:
        p_el.remove(r)
    r = runs[0]
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)


def set_cell(cell, text):
    ps = cell._tc.findall(W_P)
    for p in ps[1:]:
        cell._tc.remove(p)
    set_p(ps[0], text)


def embed_copy(path):
    """Smaller PNG for the document: at most EMBED_MAX_PX wide, 256 colours (UI frames have few colours)."""
    im = Image.open(path).convert("RGB")
    if im.width > EMBED_MAX_PX:
        im = im.resize((EMBED_MAX_PX, round(im.height * EMBED_MAX_PX / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    im.quantize(256, method=Image.Quantize.MEDIANCUT).save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else FDS
    out = sys.argv[2] if len(sys.argv) > 2 else src
    d = docx.Document(src)
    body = d.element.body

    def blocks():
        return list(body.iterchildren())

    def heading(prefix):
        for i, el in enumerate(blocks()):
            if el.tag == W_P and ptext(el).startswith(prefix) and Paragraph(el, d).style.name == "Heading 2":
                return i
        raise KeyError(prefix)

    caption_proto = next(copy.deepcopy(el) for el in blocks()
                         if el.tag == W_P and ptext(el).startswith("Figure 2.") and " — " in ptext(el))

    def picture(path):
        holder = d.add_paragraph()
        holder.alignment = 1
        pic = holder.add_run().add_picture(embed_copy(path), width=Cm(WIDTH_CM))
        if pic.height > Cm(MAX_HEIGHT_CM):
            ratio = Cm(MAX_HEIGHT_CM) / pic.height
            pic.height, pic.width = Cm(MAX_HEIGHT_CM), int(pic.width * ratio)
        body.remove(holder._p)
        return holder._p

    plan = [(num, name, sources, [(src_, tab, png) for src_ in sources for tab, png in frames(src_)])
            for num, name, sources in SCREENS]
    missing = [png for *_, shots in plan for _, _, png in shots if not os.path.exists(png)]
    if missing:
        sys.exit("PNG not exported yet (run export_diagrams.py):\n  " + "\n  ".join(missing))

    figures = 0
    for num, name, sources, shots in plan:
        h = heading(f"{num}  {name}")
        bl = blocks()
        ti = next(j for j in range(h, len(bl)) if bl[j].tag == W_TBL)
        nav = next(j for j in range(ti, len(bl)) if bl[j].tag == W_P and ptext(bl[j]) == "Navigation Context")
        spacer = bl[ti + 1]
        for el in bl[ti + 2:nav]:  # old figures and the spacers between them
            body.remove(el)
        first_pid = sources[0].split("_", 1)[0]
        new = []
        for k, (source, tab, png) in enumerate(shots, 1):
            pid, frame = tab.split(" / ", 1)
            source_name = "Shared layout" if pid == "SHELL" else by_id(pid).name
            label = frame if pid == first_pid else f"{source_name} · {frame}"
            label = label.removeprefix(f"{name} — ")  # 'Writing Feedback — Đã đạt' under 2.14 Writing Feedback
            stem = source.partition(":")[0]
            cap = copy.deepcopy(caption_proto)
            set_p(cap, f"Figure {num}-{k} — {name} — {label}. Source: wireframes/screens/{stem}.drawio, tab “{tab}”")
            new += [picture(png), cap]
        new.append(copy.deepcopy(spacer))
        for el in reversed(new):
            spacer.addnext(el)
        stems = list(dict.fromkeys(s.partition(":")[0] + ".drawio" for s in sources))
        span = f"Figure {num}-1" + (f" to {num}-{len(shots)}" if len(shots) > 1 else "")
        info = Table(bl[ti], d)
        for r in info.rows:
            if r.cells[0].text.strip() == "Mockup / Wireframe":
                set_cell(r.cells[1], f"{span} below — draw.io wireframes in docs/wireframes/screens ({', '.join(stems)}), "
                                     "one tab per state; clickable prototype: docs/wireframes/IELTSPath_Prototype.drawio")
        figures += len(shots)

    hist = Table(next(el for el in blocks() if el.tag == W_TBL and ptext(el).startswith("VersionDateAuthorDescription")), d)
    for r in hist.rows:
        if r.cells[0].text.strip() == "v0.7":
            set_cell(r.cells[3], V07_DESCRIPTION)

    used = {el.get(qn("r:embed")) for el in body.iter(qn("a:blip"))}
    for drawing_id, drawing in enumerate(body.iter(qn("wp:docPr")), 1):
        drawing.set("id", str(drawing_id))
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype.endswith("/image") and rid not in used:
            d.part.drop_rel(rid)
    d.save(out)
    print(f"{figures} figures in {len(SCREENS)} screens -> {out}")


if __name__ == "__main__":
    main()
