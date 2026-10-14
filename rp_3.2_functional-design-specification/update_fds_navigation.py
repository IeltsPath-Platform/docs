"""Patch agreed navigation and append P-54 without rebuilding customized FDS content.

Run before update_fds_figures.py. Existing unrelated sections and numbering are retained.
"""
import copy
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from update_fds_figures import FDS, ptext, set_p, set_cell
from fds_screens import SCREENS as ALL_SCREENS

SCREENS = [s for s in ALL_SCREENS if s["key"] == "SubmissionHistory"]

NAVIGATION = (
    "Design target: hover over Dashboard to open a popup directly below it with Tổng quát, Dữ liệu chi tiết, "
    "Lịch sử ôn luyện and Xếp hạng. The Dashboard has no left sidebar. The draw.io prototype simulates hover by "
    "clicking Dashboard to open P-11 / Menu Dashboard. Personal information is reached through the user menu: "
    "Thông tin cá nhân → P-80 / Hồ sơ của bạn."
)
ACTIVATION = (
    "Lịch sử kích hoạt gói shows activation-key history, not payment transactions. The history list is a design "
    "target: its listing API is not implemented. Existing subscription and key-activation APIs remain unchanged."
)
DICTIONARY = ("Dictionary lookup in scope is English to Vietnamese only (Tra Anh-Việt, API-82). The Reading glossary "
              "and the Writing Practice design (P-44) show only Tra Anh-Việt. The built Placement dictionary still has a "
              "Tra Việt-Anh tab: Vietnamese-to-English lookup is out of scope and has no API, so that tab is to be removed.")


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(FDS)
    d = docx.Document(path)
    body = d.element.body
    for p in body.iter(qn("w:p")):
        old = ptext(p)
        text = old.replace("Lịch sử thanh toán", "Lịch sử kích hoạt gói")
        text = text.replace("Lịch sử ôn luyện, Hồ sơ của tôi and Xếp hạng", "Lịch sử ôn luyện and Xếp hạng")
        text = text.replace("Lịch sử ôn luyện, Hồ sơ của tôi, Xếp hạng", "Lịch sử ôn luyện, Xếp hạng")
        if old == 'User clicks "Profile"':
            text = 'User clicks "Thông tin cá nhân" → P-80 / Hồ sơ của bạn'
        if text != old:
            set_p(p, text)

    for table in d.tables:
        for row in table.rows:
            if row.cells[0].text == "Overview-TAB-01":
                for cell, value in zip(row.cells[1:], ["Dashboard menu", "Hover popup directly below Dashboard",
                        "One active section, default Tổng quát", '"Tổng quát" · "Dữ liệu chi tiết" · "Lịch sử ôn luyện" · "Xếp hạng"']):
                    set_cell(cell, value)

    def append_note(section, text):
        for el in list(body):
            if ptext(el) == text or (section == "2.23  Practice Workspaces" and
                    ptext(el).startswith(("The dictionary remains Tra Anh-Việt", "The Reading glossary is labelled",
                                          "Dictionary lookup in scope"))):
                body.remove(el)
        blocks = list(body)
        h = next(i for i, el in enumerate(blocks) if el.tag == qn("w:p") and ptext(el).startswith(section))
        table = next(el for el in blocks[h + 1:] if el.tag == qn("w:tbl"))
        holder = d.add_paragraph(text)
        table.addprevious(holder._p)
        following = table.getnext()
        if following is None or ptext(following) or list(following.iter(qn("w:drawing"))):
            spacer = d.add_paragraph()
            table.addnext(spacer._p)

    append_note("2.7  Overview", NAVIGATION)
    append_note("2.28  Subscription & Points", ACTIVATION)
    append_note("2.23  Practice Workspaces", DICTIONARY)

    if not any(ptext(el) == "2.54  Submission History" for el in body):
        # Reuse the template's existing screen/table builders only; never run its full rebuild.
        builder = HERE / "build-fds.py"
        namespace = {"__file__": str(builder), "__name__": "fds_section_builder"}
        source = builder.read_text(encoding="utf-8").split("\nbuild_front()\n", 1)[0]
        exec(compile(source, str(builder), "exec"), namespace)
        namespace["build_screen"](54, SCREENS[0])
        anchor = next(el for el in body if ptext(el).startswith("PART 3"))
        for el in namespace["out"]:
            anchor.addprevious(copy.deepcopy(el))

    index = next(t for t in d.tables if t.rows[0].cells[0].text == "Section" and
                 t.rows[0].cells[1].text == "Screen")
    if not any(r.cells[0].text == "2.54" for r in index.rows):
        row = copy.deepcopy(index.rows[-1]._tr)
        index._tbl.append(row)
        values = ["2.54", "Submission History", SCREENS[0]["route"], SCREENS[0]["roles"],
                  "Design target (history list APIs not implemented)\nFE: placeholder route"]
        for cell, value in zip(index.rows[-1].cells, values):
            set_cell(cell, value)

    # Keep the existing API ids and add only the new screen's use of implemented detail APIs.
    for table in d.tables:
        if len(table.columns) != 7 or table.rows[0].cells[0].text != "API ID":
            continue
        for row in table.rows[1:]:
            if any(path in row.cells[1].text for path in ["/writing-submissions/{id}",
                    "/practice-attempts/{id}", "/attempts/{attemptId}/result", "/attempts/{attemptId}/placement-result"]):
                current = row.cells[4].text
                if "2.54" not in current:
                    set_cell(row.cells[4], current + (", " if current.strip("— ") else "") + "2.54")

    d.save(path)
    print(f"Updated navigation and P-54 in {path.name}")


if __name__ == "__main__":
    main()
